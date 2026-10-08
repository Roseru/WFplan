from pathlib import Path
import re, json, datetime, collections
import openpyxl, pdfplumber
from docx import Document

ROOT=Path(__file__).resolve().parent
DOWNLOADS=Path.home()/'Downloads'
clean=lambda s: re.sub(r'\s+', ' ', s or '').strip()
topics={}; lessons={}; topic=None
with pdfplumber.open(DOWNLOADS/'PROGRAM WYCHOWANIE 2024 (1).pdf') as pdf:
    for page_no in range(12,53):
        for table in pdf.pages[page_no-1].extract_tables():
            for row in table:
                if len(row)!=7: continue
                a,b,title,total,theory,practice,note=map(clean,row)
                if a.isdigit() and not b:
                    topic=a; topics[a]=title
                if b.isdigit() and topic:
                    hours=practice or theory
                    lessons[f'{topic}/{b}']={'title':title,'page':page_no,'hours':int(re.search(r'\d+',hours)[0]) if re.search(r'\d+',hours) else None,'type':'instruktorsko-metodyczne' if 'im' in hours else 'ćwiczenia' if 'ćw' in hours else 'teoretyczne'}

doc=Document(next(DOWNLOADS.glob('wytyczne_do_zajec_z_wychowania_fizycznego_semestr_zimowy_2026_2027*')))
locations={}
for p in doc.paragraphs:
    m=re.match(r'^([A-ZŻ])\s*[–-]\s*(.*?)\s*[–-]\s*obiekt (zewnętrzny|wewnętrzny)',p.text)
    if m: locations[m[1]]={'name':m[2].strip(),'outdoor':m[3]=='zewnętrzny'}
for n in [1,2,3]: locations[f'B{n}']={'name':f'Hala tartanowa, strefa B{n}'+(' — walka wręcz' if n==2 else ' — gimnastyka'),'outdoor':False}
wb=openpyxl.load_workbook(next(DOWNLOADS.glob('rozklad_zajec_zima_lato_2026-2027*')),data_only=True)
events=[]; issues=[]; semesters=[]
weekdays=['PONIEDZIAŁEK','WTOREK','ŚRODA','CZWARTEK','PIĄTEK','SOBOTA','NIEDZIELA']
for ws in wb:
    summer=ws.title.endswith('L'); sid='summer-2026' if summer else 'winter-2026'
    semesters.append({'id':sid,'label':'Letni 2025/2026 · archiwum' if summer else 'Zimowy 2026/2027','warning':'Arkusz letni zawiera plan 2025/2026. Nie jest planem na lato 2027.' if summer else '', 'start':'2026-03-01' if summer else '2026-09-28','end':'2026-06-30' if summer else '2027-01-30'})
    dates={}; block=''; day=None
    for r in range(1,ws.max_row+1):
        a=clean(str(ws.cell(r,1).value or ''))
        if a in weekdays:
            day=weekdays.index(a); dates={}; block=''
            month=3 if summer else (9 if day<=2 else 10); year=2026; last=0
            for c in range(5,22 if summer else 23):
                raw=ws.cell(r+1,c).value
                if raw is None: continue
                d=int(raw)
                if d<last:
                    month+=1
                    if month==13: month=1; year+=1
                dt=datetime.date(year,month,d)
                if dt.weekday()!=day: issues.append(f'Weekday mismatch {ws.title} {r+1}:{c} {dt}')
                dates[c]=dt.isoformat(); last=d
            continue
        if re.fullmatch(r'\d+[\\/]\d+',a): block=a.replace('\\','–').replace('/','–')
        group=clean(str(ws.cell(r,2).value or ''))
        gm=re.match(r'(\d+)pl\s*[\\/]\s*(\d+)kp(.*)',group,re.I)
        if not gm or not dates: continue
        for c,date in dates.items():
            raw=ws.cell(r,c).value
            if raw is None: continue
            raw=clean(str(raw)); match=re.fullmatch(r'([A-ZŻ]+\d?)\s*-\s*(\d+)\s*/\s*(\d+)',raw)
            event={'semester':sid,'date':date,'company':int(gm[2]),'platoon':int(gm[1]),'subgroup':gm[3].strip(),'group':group,'year':str(ws.cell(r,3).value or ''),'instructor':clean(str(ws.cell(r,4).value or '')),'block':block,'code':raw,'sheet':ws.title,'cell':ws.cell(r,c).coordinate}
            if match:
                place,t,n=match.groups(); key=f'{t}/{n}'
                event.update({'location':place,'lesson':key,'topic':t})
                if key not in lessons: issues.append('Unknown lesson '+key)
                if place not in locations: issues.append('Unknown location '+place)
            else: issues.append('Other code '+raw)
            events.append(event)
semesters.append({'id':'summer-2027','label':'Letni 2026/2027 · brak danych','warning':'W przesłanym pliku nie ma rozkładu na semestr letni 2026/2027. Zakładka „semestr L” dotyczy 2025/2026.','start':'2027-03-01','end':'2027-06-30'})
output={'updated':'2026-09-30','semesters':semesters,'topics':topics,'lessons':lessons,'locations':locations,'events':events,'sourceFiles':{'schedule':next(DOWNLOADS.glob('rozklad_zajec_zima_lato_2026-2027*')).name,'program':'PROGRAM WYCHOWANIE 2024 (1).pdf','guidelines':next(DOWNLOADS.glob('wytyczne_do_zajec_z_wychowania_fizycznego_semestr_zimowy_2026_2027*')).name}}
(ROOT/'data.json').write_text(json.dumps(output,ensure_ascii=False,separators=(',',':')),encoding='utf-8')
(ROOT/'import-report.json').write_text(json.dumps({'events':len(events),'lessons':len(lessons),'issues':dict(collections.Counter(issues))},ensure_ascii=False,indent=2),encoding='utf-8')
print(len(events),'events;',len(lessons),'lessons;',dict(collections.Counter(issues)))

